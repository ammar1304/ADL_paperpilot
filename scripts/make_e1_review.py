#!/usr/bin/env python3
"""D38 mitigation: a stratified random sample of 20 E1 test queries for a teammate to review before
the one-time test run. Writes eval/queries/test_review.xlsx (fixed seed, so the sample is reproducible)."""

import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REPO = Path(__file__).resolve().parents[1]
PER_TYPE = {"paraphrase": 7, "section": 7, "table": 6}
SEED = 38


def main() -> None:
    titles = {f"arxiv:{r['arxiv_id']}": r["title"] for r in csv.DictReader(open(REPO / "corpus/manifest.csv", encoding="utf-8"))}
    rows = [r for r in csv.DictReader(open(REPO / "eval/queries/worksheet.csv", encoding="utf-8")) if r["split"] == "test" and r["query"].strip()]
    rng = random.Random(SEED)
    sample = []
    for t, n in PER_TYPE.items():
        sample += rng.sample([r for r in rows if r["type"] == t], n)

    wb = Workbook()
    info = wb.active
    info.title = "How to review"
    for line in [
        "E1 test queries: review before the one-time test run (D38)",
        "",
        "Each query was written to find one paper of the library. For each row, answer three questions in the yellow columns:",
        "1. Realistic: would a researcher plausibly type this when looking for such a paper? (yes / no)",
        "2. Copies wording: does it reuse distinctive phrases from the context column, so that it is too easy? (yes / no)",
        "3. Right answer: is the source paper a correct answer to this query? (yes / no)",
        "Add a comment for any 'no' or 'yes' on question 2.",
        "",
        "Rules: do not look at any retrieval result first. Queries flagged here are fixed or replaced before the test run, and the report says how many.",
    ]:
        info.append([line])
    info["A1"].font = Font(bold=True, size=13)
    info.column_dimensions["A"].width = 120

    ws = wb.create_sheet("Queries")
    head = ["#", "id", "type", "query", "source paper", "what the writer saw (context)", "Realistic?", "Copies wording?", "Right answer?", "Comment"]
    ws.append(head)
    for i, r in enumerate(sample, 1):
        ws.append([i, r["id"], r["type"], r["query"], titles.get(r["relevant_papers"].split(";")[0].strip(), r["relevant_papers"]), r["context"][:700], "", "", "", ""])
    widths = [4, 22, 11, 45, 35, 70, 11, 12, 11, 30]
    for col, w in zip("ABCDEFGHIJ", widths):
        ws.column_dimensions[col].width = w
    for c in ws[1]:
        c.font = Font(bold=True)
    yellow = PatternFill("solid", fgColor="FFF4B0")
    yn = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(yn)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[6:9]:
            c.fill = yellow
            yn.add(c)
        row[9].fill = yellow
    ws.freeze_panes = "E2"
    out = REPO / "eval/queries/test_review.xlsx"
    wb.save(out)
    print(f"Wrote {out.relative_to(REPO)} ({len(sample)} queries)")


if __name__ == "__main__":
    main()
